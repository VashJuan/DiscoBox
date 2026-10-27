# SPDX-FileCopyrightText: 2026 John Cornelison for Vashon Software
# SPDX-License-Identifier: MIT

"""Cut song clips for the board from songs.xlsx (runs on the PC, not the board)

Each row of one theme (the Theme column; blank is generic) with a start time
becomes sounds/<theme>/<name>_22k.mp3: 22050 Hz mono,
32 kbps, tags stripped, to match the mixer in sound.py, with a short fade in
and a fade out. Each clip is brought to the same loudness (two-pass loudnorm,
with its gentle limiter) so no song is much louder or quieter than the rest.
Rows with no start are skipped. length is optional.
start and length are seconds, or m:ss time cells. A plain number in a cell
formatted as a time still counts as seconds. Rows with Priority above
MAX_PRIORITY are skipped too. The clips' names, in order, go in
sounds/<theme>/playlist.txt, which the board plays from.

A theme may add a sound of its own (THEME_SOUND), e.g. Halloween's
haunted-house echo; those clips are named <name>_echo_22k.mp3, and --dry cuts
them without it, so both can be compared. playlist.txt lists the last run's.

    python tools/make_clips.py                     (generic)
    python tools/make_clips.py --theme halloween   (with the echo)
    python tools/make_clips.py --theme halloween --dry
"""

import argparse
import datetime
import json
import os
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE_DIR = os.path.join(HERE, "..", "sounds_workshop")  # full-length songs, not in git
SOUNDS_DIR = os.path.join(HERE, "..", "sounds")
SONGS_XLSX = os.path.join(HERE, "songs.xlsx")

LENGTH = 60  # seconds, when the row gives none
MAX_PRIORITY = 14  # lower number = higher priority; 60 s clips through 14 fill the board
FADE_IN = 0.5  # seconds
FADE_OUT = 3  # seconds
LOUDNESS = -12  # LUFS; a little louder than streaming's -14, for outdoors
CLOSE_ENOUGH = 1  # dB; a clip further than this from LOUDNESS is cut again
PEAK = -2  # dBTP; MP3 encoding adds about 1-2 dB, so the clips don't clip
BOARD_SPACE_KB = 7000  # roughly what CIRCUITPY holds, code & libraries included

# A theme's own sound (ffmpeg filters), added before the loudness is matched
THEME_SOUND = {
    # haunted house: three echoes off far walls, fading
    "halloween": "aecho=0.8:0.8:90|230|470:0.4|0.3|0.2",
}


def seconds(value):
    """A cell as seconds: a number, or a time cell typed as m:ss"""
    if isinstance(value, datetime.datetime):  # a number shown as a date/time
        return float(to_excel(value))
    if isinstance(value, datetime.time):
        return value.hour * 3600 + value.minute * 60 + value.second
    return float(value)


def read_rows():
    """Rows of the songs sheet as dicts keyed by the header row"""
    sheet = load_workbook(SONGS_XLSX, data_only=True)["songs"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    return [dict(zip(header, row)) for row in rows]


def measure(path, start=0, length=None, before=""):
    """loudnorm's first-pass stats for a file (or part of one)"""
    target = "{}loudnorm=I={}:TP={}:LRA=11".format(before, LOUDNESS, PEAK)
    clip = ["-ss", str(start), "-t", str(length)] if length else []
    result = subprocess.run(
        ["ffmpeg", "-hide_banner", "-nostats"] + clip + [
            "-i", path, "-af", target + ":print_format=json", "-f", "null", "-",
        ],
        check=True, capture_output=True, text=True, encoding="utf-8",
    )
    log = result.stderr
    return json.loads(log[log.rindex("{"):log.rindex("}") + 1])


def loudnorm(source, start, length, extra=""):
    """loudnorm filter for one clip, from a first pass that measures it"""
    # mono & 22 kHz first, so the level measured is the level the board plays
    before = "aformat=channel_layouts=mono,aresample=22050,"
    if extra:
        before += extra + ","
    stats = measure(source, start, length, before)
    return (
        before + "loudnorm=I={}:TP={}:LRA=11".format(LOUDNESS, PEAK)
        + ":measured_I={input_i}:measured_TP={input_tp}"
        ":measured_LRA={input_lra}:measured_thresh={input_thresh}"
        ":offset={target_offset}".format(**stats)
    )


def make_clip(row, out_dir, extra="", gain=0):
    """Cut one clip with ffmpeg; returns its file name

    gain (dB) corrects a clip that loudnorm left too quiet or loud; a limiter
    then keeps its peaks under PEAK.
    """
    start = seconds(row["start"])
    length = seconds(row["length"] or LENGTH)
    source = os.path.join(SOURCE_DIR, row["source"])
    out_name = row["name"] + ("_echo" if extra else "") + "_22k.mp3"
    filters = loudnorm(source, start, length, extra)
    if gain:
        filters += ",volume={:.1f}dB,alimiter=limit={:.1f}dB".format(gain, PEAK)
    filters += ",afade=t=in:d={},afade=t=out:st={}:d={}".format(
        FADE_IN, length - FADE_OUT, FADE_OUT
    )
    subprocess.run(
        [
            "ffmpeg", "-v", "error", "-y",
            "-ss", str(start), "-t", str(length), "-i", source,
            "-af", filters, "-ac", "1", "-ar", "22050", "-b:a", "32k",
            "-map_metadata", "-1", "-id3v2_version", "0", "-write_xing", "0",
            os.path.join(out_dir, out_name),
        ],
        check=True,
    )
    return out_name


def loudness(path):
    """Integrated loudness of a finished clip, in LUFS"""
    return float(measure(path)["input_i"])


def main():
    parser = argparse.ArgumentParser(description="Cut a theme's clips for the board")
    parser.add_argument("--theme", default="generic", help="Theme column value")
    parser.add_argument(
        "--dry", action="store_true", help="leave out the theme's own sound (echo)"
    )
    args = parser.parse_args()
    extra = "" if args.dry else THEME_SOUND.get(args.theme, "")
    out_dir = os.path.join(SOUNDS_DIR, args.theme)
    os.makedirs(out_dir, exist_ok=True)

    rows = [
        row for row in read_rows()
        if (row.get("Theme") or "generic") == args.theme
        and row["start"] not in (None, "")
        and row["Priority"] is not None and row["Priority"] <= MAX_PRIORITY
    ]
    if not rows:
        sys.exit("No {} rows with start times in {}".format(args.theme, SONGS_XLSX))

    made = []
    for row in rows:
        name = make_clip(row, out_dir, extra)
        level = loudness(os.path.join(out_dir, name))
        if abs(level - LOUDNESS) > CLOSE_ENOUGH:
            name = make_clip(row, out_dir, extra, LOUDNESS - level)
            level = loudness(os.path.join(out_dir, name))
        made.append(name)
        print("made {} ({:.1f} LUFS)".format(name, level))

    total_kb = sum(
        os.path.getsize(os.path.join(out_dir, name)) for name in made
    ) // 1024
    print("\n{} clips, {} KB (board holds about {} KB in all)".format(
        len(made), total_kb, BOARD_SPACE_KB
    ))
    with open(os.path.join(out_dir, "playlist.txt"), "w", encoding="utf-8") as playlist:
        playlist.write("".join(name + "\n" for name in made))
    print("Playlist written to", os.path.join(out_dir, "playlist.txt"))
    print("To put it on the board: tools\\deploy.ps1 -Theme", args.theme)


if __name__ == "__main__":
    main()
